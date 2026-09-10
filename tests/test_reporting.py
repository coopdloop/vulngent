from __future__ import annotations

import datetime as dt

from sqlalchemy.orm import Session

from vulngent.db import repository as repo
from vulngent.db.models import Severity
from vulngent.report_data import collect_report_data
from vulngent.reporting import _pdf_safe, render_report


def _seed(session: Session) -> None:
    asset = repo.get_or_create_asset(session, "svc-a", criticality="critical")
    now = dt.datetime.now(dt.timezone.utc)
    overdue = repo.create_vulnerability(
        session, external_id="CVE-OLD", title="Overdue finding", severity=Severity.CRITICAL, cvss_score=9.8,
        asset=asset, discovered_at=now - dt.timedelta(days=30),
    )
    repo.set_priority(session, overdue)
    fresh = repo.create_vulnerability(
        session, external_id="CVE-NEW", title="Fresh finding", severity=Severity.HIGH, cvss_score=7.5,
        asset=asset, discovered_at=now,
    )
    repo.set_priority(session, fresh)


def test_render_markdown_includes_overdue_summary(session: Session) -> None:
    _seed(session)
    data = collect_report_data(session)

    md = render_report(data, "md")

    assert "Overdue: 1" in md
    assert "## Past SLA (overdue)" in md
    assert "CVE-OLD" in md
    assert "[OVERDUE" in md


def test_render_markdown_empty_ledger_has_none_placeholder(session: Session) -> None:
    data = collect_report_data(session)

    md = render_report(data, "md")

    assert "Overdue: 0" in md
    assert "## Past SLA (overdue)\n- None" in md


def test_pdf_safe_transliterates_common_unicode_punctuation() -> None:
    assert _pdf_safe("em\u2014dash") == "em-dash"
    assert _pdf_safe("curly \u2018quotes\u2019") == "curly 'quotes'"
    assert _pdf_safe("bullet \u2022 point") == "bullet - point"


def test_pdf_safe_falls_back_instead_of_raising_on_unsupported_glyphs() -> None:
    # CJK has no latin-1 representation; must not raise, must return latin-1-safe text.
    safe = _pdf_safe("\u4f60\u597d")
    safe.encode("latin-1")  # raises if not actually safe


def test_render_report_pdf_produces_valid_pdf_bytes(session: Session) -> None:
    _seed(session)
    data = collect_report_data(session)

    rendered = render_report(data, "pdf")

    assert isinstance(rendered, bytes)
    assert rendered.startswith(b"%PDF-")
    assert len(rendered) > 1000


def test_render_report_pdf_handles_empty_ledger(session: Session) -> None:
    data = collect_report_data(session)

    rendered = render_report(data, "pdf")

    assert isinstance(rendered, bytes)
    assert rendered.startswith(b"%PDF-")


def test_render_report_docx_produces_valid_docx_bytes(session: Session) -> None:
    import io

    from docx import Document

    _seed(session)
    data = collect_report_data(session)

    rendered = render_report(data, "docx")

    assert isinstance(rendered, bytes)
    doc = Document(io.BytesIO(rendered))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    # Title comes from REPORT_TITLE; assert against the collected branding, not a
    # literal, so a customized local .env can't break the suite.
    assert data.branding.title in full_text

    # KPI cards + findings + overdue tables; no commitments were seeded so that
    # section renders as a paragraph, not a table.
    assert len(doc.tables) == 3


def test_render_report_rejects_unsupported_format(session: Session) -> None:
    import pytest

    data = collect_report_data(session)
    with pytest.raises(ValueError, match="Unsupported report format"):
        render_report(data, "pptx")
